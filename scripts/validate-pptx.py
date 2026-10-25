#!/usr/bin/env python3
"""Validate a SlydeCode-produced .pptx with python-pptx.

This is the independent check for `crates/slyde-pptx`: it opens the deck with a
real OOXML consumer (not our own writer) and asserts the package is what the
export contract promises.

Usage:
    python scripts/validate-pptx.py target/tmp/pptx-check.pptx --expect-slides 3
    python scripts/validate-pptx.py deck.pptx --expect-slides 3 --expect-size 1280x720

Exit code 0 means every requested check passed; any failure prints a FAIL line
and returns 1.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_PIXEL = 9525


def parse_size(value: str) -> tuple[int, int]:
    """Parse a WxH pixel size such as ``1280x720``."""
    try:
        width, height = value.lower().split("x", 1)
        return int(width), int(height)
    except ValueError as error:
        raise argparse.ArgumentTypeError(f"expected WxH, got {value!r}") from error


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate a .pptx deck.")
    parser.add_argument("path", type=Path, help="path to the .pptx file")
    parser.add_argument(
        "--expect-slides",
        type=int,
        required=True,
        help="expected number of slides",
    )
    parser.add_argument(
        "--expect-size",
        type=parse_size,
        default=None,
        metavar="WxH",
        help="expected slide size in CSS pixels at 96 DPI (e.g. 1280x720)",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    failures: list[str] = []

    try:
        presentation = Presentation(str(args.path))
    except Exception as error:  # noqa: BLE001 - any open failure is a FAIL
        print(f"FAIL could not open {args.path}: {error!r}", file=sys.stderr)
        return 1

    slides = list(presentation.slides)
    width_emu = presentation.slide_width
    height_emu = presentation.slide_height
    print(
        f"opened {args.path}: {len(slides)} slides, "
        f"{width_emu}x{height_emu} EMU"
    )

    if len(slides) != args.expect_slides:
        failures.append(
            f"expected {args.expect_slides} slides, found {len(slides)}"
        )

    if args.expect_size is not None:
        expected_width = args.expect_size[0] * EMU_PER_PIXEL
        expected_height = args.expect_size[1] * EMU_PER_PIXEL
        print(
            f"expected size {args.expect_size[0]}x{args.expect_size[1]} px = "
            f"{expected_width}x{expected_height} EMU"
        )
        if width_emu != expected_width:
            failures.append(f"slide width is {width_emu} EMU, expected {expected_width}")
        if height_emu != expected_height:
            failures.append(
                f"slide height is {height_emu} EMU, expected {expected_height}"
            )

    for number, slide in enumerate(slides, start=1):
        pictures = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        if len(pictures) != 1:
            failures.append(f"slide {number}: expected 1 picture, found {len(pictures)}")
            continue

        picture = pictures[0]
        bounds = (picture.left, picture.top, picture.width, picture.height)
        print(f"slide {number}: one picture at {bounds}, {slide.name!r}")
        if bounds != (0, 0, width_emu, height_emu):
            failures.append(
                f"slide {number}: picture does not cover the slide, "
                f"frame is {bounds}, slide is {(0, 0, width_emu, height_emu)}"
            )
        if picture.image.content_type != "image/png":
            failures.append(
                f"slide {number}: image content type is "
                f"{picture.image.content_type!r}, expected 'image/png'"
            )
        if args.expect_size is not None and picture.image.size != args.expect_size:
            failures.append(
                f"slide {number}: PNG is {picture.image.size[0]}x"
                f"{picture.image.size[1]} px, expected {args.expect_size[0]}x"
                f"{args.expect_size[1]}"
            )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            first_line = notes.splitlines()[0] if notes else ""
            print(f"slide {number}: notes {len(notes)} chars, first line {first_line!r}")

    if failures:
        for failure in failures:
            print(f"FAIL {failure}", file=sys.stderr)
        print(f"FAILED {args.path}: {len(failures)} check(s) failed", file=sys.stderr)
        return 1

    print(f"OK {args.path}: {len(slides)} slides, one full-bleed PNG each")
    return 0


if __name__ == "__main__":
    sys.exit(main())
